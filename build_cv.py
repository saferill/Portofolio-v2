"""Create bilingual, text-based CV PDFs and editable DOCX files from portfolio data.
Install requirements-documents.txt before running. No invented education dates or metrics.
"""
from pathlib import Path
from xml.sax.saxutils import escape
import json
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
OUT=ROOT/'site/documents';OUT.mkdir(exist_ok=True)
fontdir=ROOT/'source/fonts'
for name,file in [('CV','DejaVuSans.ttf'),('CVBold','DejaVuSans-Bold.ttf')]:
 pdfmetrics.registerFont(TTFont(name,str(fontdir/file)))
pdfmetrics.registerFontFamily('CV',normal='CV',bold='CVBold',italic='CV',boldItalic='CVBold')
ink=colors.HexColor('#17212b');accent=colors.HexColor('#315569');muted=colors.HexColor('#53616a')
styles={
 'name':ParagraphStyle('name',fontName='CVBold',fontSize=21,leading=26,textColor=ink,spaceAfter=4),
 'tagline':ParagraphStyle('tagline',fontName='CV',fontSize=10,leading=14,textColor=accent,spaceAfter=6),
 'contact':ParagraphStyle('contact',fontName='CV',fontSize=9,leading=12.5,textColor=muted),
 'heading':ParagraphStyle('heading',fontName='CVBold',fontSize=10,leading=14,textColor=accent,spaceBefore=12,spaceAfter=5),
 'body':ParagraphStyle('body',fontName='CV',fontSize=10,leading=14.5,textColor=ink,spaceAfter=3),
 'sub':ParagraphStyle('sub',fontName='CVBold',fontSize=10,leading=14,textColor=ink,spaceAfter=2),
 'meta':ParagraphStyle('meta',fontName='CV',fontSize=9,leading=12.5,textColor=muted,spaceAfter=3),
 'bullet':ParagraphStyle('bullet',fontName='CV',fontSize=10,leading=14.5,textColor=ink,leftIndent=10,firstLineIndent=-8,spaceAfter=2)
}
content=json.loads((ROOT/'site/data/cv-content.json').read_text())

def doc_hyperlink(paragraph,label,url):
 rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
 run=OxmlElement('w:r');prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'315569');prop.append(color);size=OxmlElement('w:sz');size.set(qn('w:val'),'17');prop.append(size);run.append(prop);text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);paragraph._p.append(link)

for lang in ['id','en']:
 p=json.loads((ROOT/'site/data'/('profile-id.json' if lang=='id' else 'profile.json')).read_text());c=content[lang]
 blocks=[]
 def add(kind,text,link=None):blocks.append((kind,text,link))
 add('name',p['name']);add('tagline',c['headline'])
 add('contact',p['location']+'  |  +62 851-4300-1281')
 add('contact',p['email'],'mailto:'+p['email'])
 add('heading',c['summaryTitle'].upper());add('body',c['summary'])
 for experience,title in zip(p['experience'],[c['experienceTitle'],c['organizationTitle']]):
  add('heading',title.upper());add('sub',experience['role']);add('body',experience['organization']);add('meta',experience['period'])
  for detail in experience['details']:add('bullet',detail)
 add('heading',c['educationTitle'].upper());add('sub',p['university']);add('body',c['current']);add('sub',p['school']);add('body',c['previous'])
 add('heading',c['skillsTitle'].upper());add('body',' · '.join(p['skills']))
 add('heading',c['projectsTitle'].upper());add('sub','Word Academic');add('body',c['word']);add('contact','lynk.id/wordacademic','https://lynk.id/wordacademic')
 add('sub',c['webTitle']);add('body',c['web'])
 add('contact','mailtemppro.netlify.app  |  codamusic.vercel.app  |  safevideos.vercel.app')
 pdf=[]
 for kind,text,url in blocks:
  markup=escape(text)
  if url:markup=f'<link href="{escape(url)}" color="#315569">{markup}</link>'
  if kind=='bullet':markup='• '+markup
  pdf.append(Paragraph(markup,styles[kind]))
  if kind=='heading':pass
 SimpleDocTemplate(str(OUT/f'syafril-cv-{lang}.pdf'),pagesize=A4,leftMargin=44,rightMargin=44,topMargin=34,bottomMargin=34,title=p['name']+' — CV ('+lang.upper()+')',author=p['name']).build(pdf)
 doc=Document();section=doc.sections[0];section.page_width=Pt(A4[0]);section.page_height=Pt(A4[1]);section.top_margin=Pt(34);section.bottom_margin=Pt(34);section.left_margin=Pt(44);section.right_margin=Pt(44)
 normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.font.color.rgb=RGBColor.from_string('17212B');normal.paragraph_format.space_after=Pt(3);normal.paragraph_format.line_spacing=1.05
 doc.core_properties.title=p['name']+' — CV ('+lang.upper()+')';doc.core_properties.author=p['name']
 for kind,text,url in blocks:
  par=doc.add_paragraph();fmt=par.paragraph_format
  if kind=='heading':
   par.style=doc.styles['Heading 1'];fmt.space_before=Pt(11);fmt.space_after=Pt(4);fmt.keep_with_next=True
  elif kind in ['name','tagline','sub']:fmt.keep_with_next=True
  if kind=='bullet':fmt.left_indent=Pt(10);fmt.first_line_indent=Pt(-8);text='• '+text
  if url:doc_hyperlink(par,text,url)
  else:
   run=par.add_run(text);run.font.name='Calibri';run.font.size=Pt({'name':22,'tagline':10.5,'contact':9,'heading':10,'sub':10,'meta':9}.get(kind,10));run.bold=kind in ['name','heading','sub'];run.font.color.rgb=RGBColor.from_string('315569' if kind in ['tagline','heading'] else '53616A' if kind in ['contact','meta'] else '17212B')
 doc.save(OUT/f'syafril-cv-{lang}.docx')
 print('Created',lang.upper(),'PDF and DOCX')
